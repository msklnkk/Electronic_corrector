# backend/src/project/gost_checker/parser.py
from collections import Counter
from pathlib import Path
from typing import Dict, Any, Optional
import io
import re

from docx import Document as DocxDocument
import pdfplumber


STOP_PATTERNS = [
    r'^(import|from|async|def|class|return|if|else|try|except|const|export|interface)\b',
    r'[{}\[\]();=><]',
    r'^\s*["\']',
    r'@\w+',
    r'\.\w+\(',
]

def is_valid_heading(text: str) -> bool:
    for pattern in STOP_PATTERNS:
        if re.search(pattern, text, re.IGNORECASE):
            return False
    if len(text) > 100:
        return False
    return True

def normalize_line_spacing(value) -> float:
    if value is None:
        return None
    if value > 100:
        pt = value / 12700
        return round(pt / 14, 1)
    return round(value, 1)


async def extract_document_data(
        file_path: Optional[str] = None,
        file_content: Optional[bytes] = None,
        filename: Optional[str] = None
) -> Dict[str, Any]:
    # Извлекает данные из документа для проверки по ГОСТ.
    #
    # Поддерживает два режима работы:
    # 1. file_path — для текущего FastAPI (как было раньше)
    # 2. file_content + filename — для gRPC (байты в памяти)
    # Определяем расширение и источник данных
    if file_content is not None:
        if filename is None:
            raise ValueError("При передаче file_content обязательно укажите filename для определения типа файла")
        ext = Path(filename).suffix.lower()
        file_stream = io.BytesIO(file_content)
    elif file_path is not None:
        path = Path(file_path)
        ext = path.suffix.lower()
        file_stream = None
    else:
        raise ValueError("Необходимо передать либо file_path, либо (file_content + filename)")

    data: Dict[str, Any] = {
        "required_elements": [],
        "font_settings": {"font_family": None, "font_size": None, "line_spacing": None},
        "page_margins": {"left": None, "right": None, "top": None, "bottom": None},
        "paragraph_indent": None,
        "introduction_text": "",
        "full_text": "",
        "page_count": 0,
        "word_count": 0,
    }

    try:
        if ext == ".docx":
            # ====================== DOCX ======================
            if file_content is not None:
                doc = DocxDocument(file_stream)
            else:
                doc = DocxDocument(file_path)

            fonts = []
            sizes = []
            line_spacings = []
            indents = []
            has_chapters = False
            intro_started = False
            intro_text = []

            for para_num, para in enumerate(doc.paragraphs):
                text = para.text.strip()
                if not text:
                    continue

                data["full_text"] += text + " "
                data["word_count"] += len(text.split())

                lower_text = text.lower()

                # Определение заголовков
                is_heading = (
                        para.style.name.startswith("Heading") or
                        any(getattr(run, 'bold', False) for run in para.runs) or
                        (len(text) < 50 and text.isupper()) or
                        re.match(r"^\d+\.?\s", lower_text)
                )

                if is_heading:
                    cleaned = re.sub(r'^\d+\.?\s*', '', lower_text).strip()
                    if cleaned and cleaned not in data["required_elements"] and is_valid_heading(cleaned):
                        data["required_elements"].append(cleaned)
                    if re.match(r'^\d+\.', lower_text) or para.style.name.startswith("Heading 1"):
                        has_chapters = True

                # Титульный лист
                if para_num < 50 and any(keyword in lower_text for keyword in [
                    "курсовая работа", "дипломная работа", "выпускная квалификационная работа",
                    "студент", "преподаватель", "университет", "факультет"
                ]):
                    if "титульный лист" not in data["required_elements"]:
                        data["required_elements"].append("титульный лист")

                # Введение
                if "введение" in lower_text:
                    intro_started = True
                if intro_started:
                    intro_text.append(text)

                # Форматирование
                CODE_FONTS = {'courier new', 'courier', 'consolas', 'lucida console', 'monaco', 'georgia'}

                # Шрифт параграфа из стиля
                para_font_name = None
                para_font_size = None

                # Сначала пробуем стиль параграфа
                if para.style and para.style.font:
                    para_font_name = para.style.font.name
                    para_font_size = para.style.font.size.pt if para.style.font.size else None

                # Потом runs — они перекрывают стиль
                if para.runs:
                    for run in para.runs:
                        font_name = getattr(run.font, 'name', None) or para_font_name
                        font_size = run.font.size.pt if getattr(run.font, 'size', None) else para_font_size

                        # Пропускаем кодовые шрифты
                        if font_name and font_name.lower() in CODE_FONTS:
                            continue

                        if font_name:
                            fonts.append(font_name)
                        if font_size and font_size >= 10:
                            sizes.append(font_size)

                if para.paragraph_format.line_spacing is not None:
                    line_spacings.append(para.paragraph_format.line_spacing)

                if para.paragraph_format.first_line_indent is not None:
                    indent_cm = para.paragraph_format.first_line_indent.cm
                    if indent_cm > 0:
                        indents.append(indent_cm)

            data["introduction_text"] = " ".join(intro_text).strip()

            # Агрегация значений
            data["font_settings"]["font_family"] = (
                max(Counter(fonts).items(), key=lambda x: x[1])[0] if fonts else None
            )

            sizes_filtered = [s for s in sizes if s >= 10]
            data["font_settings"]["font_size"] = (
                round(max(Counter(sizes_filtered).items(), key=lambda x: x[1])[0])
                if sizes_filtered else None
            )

            raw_spacing = (
                max(Counter(line_spacings).items(), key=lambda x: x[1])[0]
                if line_spacings else None
            )
            data["font_settings"]["line_spacing"] = normalize_line_spacing(raw_spacing)
            data["paragraph_indent"] = round(
                max(Counter(indents).items(), key=lambda x: x[1])[0], 2) if indents else None

            # Поля страницы
            if doc.sections:
                s = doc.sections[0]
                data["page_margins"] = {
                    "left": round(s.left_margin.mm, 1) if s.left_margin else None,
                    "right": round(s.right_margin.mm, 1) if s.right_margin else None,
                    "top": round(s.top_margin.mm, 1) if s.top_margin else None,
                    "bottom": round(s.bottom_margin.mm, 1) if s.bottom_margin else None,
                }

            if has_chapters:
                if "основная часть" not in data["required_elements"]:
                    data["required_elements"].append("основная часть")

        elif ext == ".pdf":
            # ====================== PDF ======================
            if file_content is not None:
                pdf = pdfplumber.open(file_stream)
            else:
                pdf = pdfplumber.open(file_path)

            with pdf:
                data["page_count"] = len(pdf.pages)
                fonts = []
                sizes = []
                line_y_diffs = []
                body_indents = []
                all_x0, all_x1, all_y0, all_y1 = [], [], [], []

                has_chapters = False
                full_text_lines = []

                for page_num, page in enumerate(pdf.pages, 1):
                    text = page.extract_text() or ""
                    data["full_text"] += text + " "
                    data["word_count"] += len(re.findall(r'\w+', text))

                    # Структура документа
                    char_sizes = [char.get('size', 0) for char in page.chars if char.get('size', 0) > 0]
                    avg_size = sum(char_sizes) / len(char_sizes) if char_sizes else 12
                    lines = [l.strip() for l in text.split("\n") if l.strip()]
                    full_text_lines.extend(lines)

                    for line in lines:
                        lower_line = line.lower()
                        pattern = r"^(содержание|оглавление|определения|обозначения и сокращения|введение|заключение|список (использованных|литературы|источников)|приложения|\d+\.?\s?[\w]+|глава \d+)"
                        if re.match(pattern, lower_line) or any(
                                char.get('size', 0) > avg_size + 1
                                for char in page.chars
                                if char.get('text') in line
                        ):
                            cleaned = re.sub(r'^\d+\.?\s*', '', lower_line).strip()
                            if cleaned and cleaned not in data["required_elements"]:
                                data["required_elements"].append(cleaned)
                            if re.match(r'^\d+\.', line):
                                has_chapters = True

                    # Титульный лист
                    if page_num == 1:
                        page_text_lower = text.lower()
                        title_keywords = [
                            "курсовая работа", "дипломная работа", "выпускная квалификационная",
                            "студент", "преподаватель", "университет", "факультет",
                            "институт", "кафедра", "направление", "специальность"
                        ]
                        matches = sum(1 for kw in title_keywords if kw in page_text_lower)
                        if matches >= 2:  # минимум 2 совпадения = титульный лист
                            if "титульный лист" not in data["required_elements"]:
                                data["required_elements"].append("титульный лист")

                    # Сбор шрифтов и размеров
                    page_height = page.height
                    for char in page.chars:
                        if 'fontname' in char:
                            font_name = re.sub(r'^[A-Z]{6}\+', '', char['fontname'])
                            font_name = re.sub(r'-(Regular|Bold|Italic|BoldItalic)$', '', font_name)
                            fonts.append(font_name)
                        if 'size' in char and char['size'] > 0:
                            sizes.append(char['size'])

                        char_y1 = char.get('y1', 0)
                        char_y0 = char.get('y0', 0)

                        is_header_footer = (
                                char_y1 > page_height * 0.95 or  # верхний колонтитул
                                char_y0 < page_height * 0.05  # нижний колонтитул / номер страницы
                        )

                        if not is_header_footer and char.get('text', '').strip():
                            if 'x0' in char:
                                all_x0.append(char['x0'])
                            if 'x1' in char:
                                all_x1.append(char['x1'])
                            if 'y0' in char:
                                all_y0.append(char['y0'])
                            if 'y1' in char:
                                all_y1.append(char['y1'])

                    # Line spacing и indent
                    y_positions = sorted(set(char['y0'] for char in page.chars if 'y0' in char), reverse=True)
                    for i in range(len(y_positions) - 1):
                        diff = y_positions[i] - y_positions[i + 1]
                        if avg_size * 0.5 < diff < avg_size * 3:
                            line_y_diffs.append(diff)

                            # Абзацный отступ - только основной текст.
                            # Ищем строки, где x0 заметно больше минимального x0 страницы
                            page_chars_body = [
                                c for c in page.chars
                                if c.get('text', '').strip()
                                   and c.get('size', 0) <= avg_size + 1  # не заголовки
                                   and page_height * 0.05 < c.get('y0', 0) < page_height * 0.95
                            ]

                            if page_chars_body:
                                min_x0_page = min(c['x0'] for c in page_chars_body)

                                # Группируем символы по строкам
                                lines_dict: Dict[float, list] = {}
                                for c in page_chars_body:
                                    y_key = round(c['y1'], 0)
                                    if y_key not in lines_dict:
                                        lines_dict[y_key] = []
                                    lines_dict[y_key].append(c)

                                for y_key, chars_in_line in lines_dict.items():
                                    first_x = min(c['x0'] for c in chars_in_line)
                                    indent_pt = first_x - min_x0_page
                                    # Абзацный отступ ~1.25см = ~35.4pt
                                    # Берем только строки с заметным отступом (15–60pt)
                                    if 15 < indent_pt < 60:
                                        body_indents.append(indent_pt / 28.346)  # pt → cm

                data["introduction_text"] = extract_introduction_from_text(data["full_text"])

                # Агрегация
                data["font_settings"]["font_family"] = (
                    max(Counter(fonts).items(), key=lambda x: x[1])[0] if fonts else None
                )
                data["font_settings"]["font_size"] = (
                    round(max(Counter(sizes).items(), key=lambda x: x[1])[0]) if sizes else None
                )

                avg_diff = sum(line_y_diffs) / len(line_y_diffs) if line_y_diffs else None
                avg_size_final = data["font_settings"]["font_size"] or 14
                data["font_settings"]["line_spacing"] = (
                    round(avg_diff / avg_size_final, 1) if avg_diff else None
                )

                if body_indents:
                    body_indents_sorted = sorted(body_indents)
                    mid = len(body_indents_sorted) // 2
                    data["paragraph_indent"] = round(body_indents_sorted[mid], 2)
                else:
                    data["paragraph_indent"] = None

                # Margins
                pt_to_mm = 0.3528
                if all_x0 and all_x1 and all_y0 and all_y1:
                    all_x0_s = sorted(all_x0)
                    all_x1_s = sorted(all_x1)
                    all_y0_s = sorted(all_y0)
                    all_y1_s = sorted(all_y1)

                    p5 = lambda lst: lst[max(0, len(lst) // 20)]
                    p95 = lambda lst: lst[min(len(lst) - 1, len(lst) * 19 // 20)]

                    min_x = p5(all_x0_s)
                    max_x = p95(all_x1_s)
                    min_y = p5(all_y0_s)
                    max_y = p95(all_y1_s)

                    width = pdf.pages[0].width if pdf.pages else 595
                    height = pdf.pages[0].height if pdf.pages else 842
                    data["page_margins"] = {
                        "left": round(min_x * pt_to_mm, 1),
                        "right": round((width - max_x) * pt_to_mm, 1),
                        "top": round((height - max_y) * pt_to_mm, 1),
                        "bottom": round(min_y * pt_to_mm, 1),
                    }

                if has_chapters:
                    if "основная часть" not in data["required_elements"]:
                        data["required_elements"].append("основная часть")

        else:
            raise ValueError(f"Неподдерживаемый формат: {ext}")

    except Exception as e:
        raise RuntimeError(f"Ошибка при парсинге документа ({ext}): {str(e)}") from e

    return data




def extract_introduction_from_text(full_text: str) -> str:
    text_lower = full_text.lower()

    all_matches = list(re.finditer(r'\bвведение\b', text_lower))

    if not all_matches:
        print("=== ВВЕДЕНИЕ НЕ НАЙДЕНО ===")
        return full_text[:2000]

    intro_start_match = None
    for match in all_matches:
        pos = match.end()
        snippet = full_text[pos: pos + 100]
        dot_count = snippet.count('.')
        if dot_count < 10:
            intro_start_match = match
            break

    if intro_start_match is None:
        intro_start_match = all_matches[-1]

    start_pos = intro_start_match.end()
    search_text = full_text[start_pos:]

    # Ищем конец введения вручную по строкам — надёжнее регулярок
    lines = search_text.split('\n')
    end_pos = len(search_text)

    for i, line in enumerate(lines):
        stripped = line.strip()
        stripped_lower = stripped.lower()

        # Пропускаем пустые строки
        if not stripped:
            continue

        # Явные маркеры конца введения
        explicit_end = (
                stripped_lower.startswith('заключение') or
                stripped_lower.startswith('conclusion') or
                re.match(r'^список\s+(использованных|литературы|источников)', stripped_lower) or
                stripped_lower.startswith('основная часть')
        )
        if explicit_end:
            end_pos = sum(len(l) + 1 for l in lines[:i])
            break

        if re.match(r'^\d+\s+\d+\s+\S', stripped):
            end_pos = sum(len(l) + 1 for l in lines[:i])
            break

        # Заголовок главы — короткая строка (до 60 символов)
        # начинается с цифры, только заглавные или Первая заглавная
        # НЕ является пунктом списка (пункты длиннее и содержат глаголы)
        is_chapter_heading = (
                re.match(r'^\d+\.?\s+[А-ЯЁA-Z]', stripped) and
                len(stripped) < 60 and  # короткая строка
                not re.search(r'\.$', stripped) and  # не заканчивается точкой (пункты списка)
                not re.search(r'\[', stripped)  # нет ссылок типа [1]
        )
        if is_chapter_heading:
            end_pos = sum(len(l) + 1 for l in lines[:i])
            break

    intro_text = search_text[:end_pos]

    return intro_text.strip()
